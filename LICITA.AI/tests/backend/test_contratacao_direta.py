from pathlib import Path
import unittest

import config
from fila import _eh_contratacao_direta


class ContratacaoDiretaTest(unittest.TestCase):
    def test_modelo_de_contratacao_direta_esta_configurado(self):
        self.assertIn("Contratação Direta - BASE.docx", config.BASE_FILES_CONTRATACAO_DIRETA)
        self.assertTrue((Path(config.PASTA_MODELOS) / "Contratação Direta - BASE.docx").is_file())

    def test_apenas_dispensa_eh_contratacao_direta(self):
        self.assertTrue(_eh_contratacao_direta({"{{MODALIDADE}}": "DISPENSA_EMAIL"}))
        self.assertTrue(_eh_contratacao_direta({"{{MODALIDADE}}": "DISPENSA_BLL"}))
        self.assertFalse(_eh_contratacao_direta({"{{MODALIDADE}}": "PREGAO_ELETRONICO"}))

    def test_modelo_tem_placeholders_de_contratacao_direta(self):
        from docx import Document

        caminho = Path(config.PASTA_MODELOS) / "Contratação Direta - BASE.docx"
        doc = Document(caminho)
        textos = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    textos.extend(p.text for p in cell.paragraphs)
        conteudo = "\n".join(textos)

        for placeholder in [
            "{{FUNDAMENTO_CONTRATACAO_DIRETA}}",
            "{{JUSTIFICATIVA_CONTRATACAO_DIRETA}}",
            "{{INSTRUCAO_ART72}}",
            "{{CRITERIOS_ESCOLHA_FORNECEDOR}}",
            "{{CONDICOES_CONTRATACAO}}",
            "{{PUBLICIDADE_TRANSPARENCIA}}",
            "{{CONCLUSAO_CONTRATACAO_DIRETA}}",
            "{{ITENS}}",
            "{{FORNECEDOR}}",
            "{{CNPJ_FORNECEDOR}}",
            "{{VALOR_CONTRATADO}}",
        ]:
            self.assertIn(placeholder, conteudo)


if __name__ == "__main__":
    unittest.main()
