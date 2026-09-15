from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document

TEST_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(TEST_ROOT))

from processador_docx import replace_text_in_paragraph


class ProcessadorDocxTests(unittest.TestCase):
    def test_tabela_aceita_linhas_com_chaves_diferentes(self):
        doc = Document()
        paragraph = doc.add_paragraph("{{ITENS}}")
        dados = [
            {"Item": 1, "Descrição": "Caneta"},
            {"Item": 2, "Descrição": "Papel", "Observação": "Entrega imediata"},
        ]

        replace_text_in_paragraph(
            paragraph,
            {"{{ITENS}}": "__TABLE__" + json.dumps(dados, ensure_ascii=False)},
        )

        with tempfile.TemporaryDirectory() as temp_dir:
            caminho = Path(temp_dir) / "teste.docx"
            doc.save(caminho)
            reaberto = Document(caminho)

        self.assertEqual(len(reaberto.tables), 1)
        tabela = reaberto.tables[0]
        self.assertEqual(len(tabela.columns), 3)
        self.assertEqual(tabela.cell(0, 2).text, "Observação")
        self.assertEqual(tabela.cell(1, 2).text, "")
        self.assertEqual(tabela.cell(2, 2).text, "Entrega imediata")

    def test_placeholder_dividido_em_runs_preserva_formatacao_ao_redor(self):
        doc = Document()
        paragraph = doc.add_paragraph()

        prefixo = paragraph.add_run("Introdução ")
        inicio_chave = paragraph.add_run("{{CH")
        inicio_chave.bold = True
        fim_chave = paragraph.add_run("AVE}}")
        fim_chave.italic = True
        sufixo = paragraph.add_run(" final")
        sufixo.underline = True

        replace_text_in_paragraph(paragraph, {"{{CHAVE}}": "resultado"})

        self.assertEqual(paragraph.text, "Introdução resultado final")
        self.assertEqual(len(paragraph.runs), 3)
        self.assertFalse(paragraph.runs[0].bold)
        self.assertFalse(paragraph.runs[0].italic)
        self.assertTrue(paragraph.runs[1].bold)
        self.assertFalse(paragraph.runs[1].italic)
        self.assertFalse(paragraph.runs[2].bold)
        self.assertFalse(paragraph.runs[2].italic)
        self.assertTrue(paragraph.runs[2].underline)

        with tempfile.TemporaryDirectory() as temp_dir:
            caminho = Path(temp_dir) / "teste_formatacao.docx"
            doc.save(caminho)
            reaberto = Document(caminho)
            runs = reaberto.paragraphs[0].runs

        self.assertEqual(reaberto.paragraphs[0].text, "Introdução resultado final")
        self.assertTrue(runs[1].bold)
        self.assertTrue(runs[2].underline)

    def test_placeholder_dividido_em_tres_runs_preserva_estilos_dos_trechos(self):
        doc = Document()
        paragraph = doc.add_paragraph()

        primeira = paragraph.add_run("A ")
        primeira.font.name = "Arial"
        parte1 = paragraph.add_run("{{OB")
        parte1.font.name = "Calibri"
        parte1.bold = True
        parte2 = paragraph.add_run("JET")
        parte2.italic = True
        parte3 = paragraph.add_run("O}}")
        parte3.underline = True
        ultima = paragraph.add_run(" B")
        ultima.font.name = "Times New Roman"

        replace_text_in_paragraph(paragraph, {"{{OBJETO}}": "MATERIAL"})

        self.assertEqual(paragraph.text, "A MATERIAL B")
        self.assertEqual(len(paragraph.runs), 3)
        self.assertEqual(paragraph.runs[0].font.name, "Arial")
        self.assertEqual(paragraph.runs[1].font.name, "Calibri")
        self.assertTrue(paragraph.runs[1].bold)
        self.assertFalse(paragraph.runs[1].italic)
        self.assertEqual(paragraph.runs[2].font.name, "Times New Roman")


if __name__ == "__main__":
    unittest.main()
